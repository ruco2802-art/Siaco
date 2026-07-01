# -*- coding: utf-8 -*-
"""Generador de documentos formales de oferta — SIACO v3.1"""
import io
import json
import zipfile
from datetime import date
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, HTTPException, Header
from fastapi.responses import Response
from pydantic import BaseModel

router = APIRouter(tags=["generador_oferta"])

FORMULARIOS_DIR = Path("biblioteca_normativa/formularios")

MESES_ES = ["enero","febrero","marzo","abril","mayo","junio",
            "julio","agosto","septiembre","octubre","noviembre","diciembre"]


# ── Utilidades ────────────────────────────────────────────────────────────────

def numero_a_letras(valor: float) -> str:
    """
    Convierte un valor monetario a texto en español para documentos legales.
    Ej: 197_500_000 → 'CIENTO NOVENTA Y SIETE MILLONES QUINIENTOS MIL PESOS M/CTE'
    """
    n = int(round(abs(valor)))
    if n == 0:
        return "CERO PESOS M/CTE"

    ONES  = ["", "UN", "DOS", "TRES", "CUATRO", "CINCO", "SEIS", "SIETE", "OCHO", "NUEVE",
             "DIEZ", "ONCE", "DOCE", "TRECE", "CATORCE", "QUINCE", "DIECISÉIS",
             "DIECISIETE", "DIECIOCHO", "DIECINUEVE"]
    TENS  = ["", "DIEZ", "VEINTE", "TREINTA", "CUARENTA", "CINCUENTA",
             "SESENTA", "SETENTA", "OCHENTA", "NOVENTA"]
    HUNDS = ["", "CIENTO", "DOSCIENTOS", "TRESCIENTOS", "CUATROCIENTOS", "QUINIENTOS",
             "SEISCIENTOS", "SETECIENTOS", "OCHOCIENTOS", "NOVECIENTOS"]
    VEINTI = {21: "VEINTIÚN", 22: "VEINTIDÓS", 23: "VEINTITRÉS", 24: "VEINTICUATRO",
              25: "VEINTICINCO", 26: "VEINTISÉIS", 27: "VEINTISIETE",
              28: "VEINTIOCHO", 29: "VEINTINUEVE"}

    def _grupo(m: int) -> str:
        if m == 0:   return ""
        if m == 100: return "CIEN"
        parts = []
        c = m // 100; r = m % 100
        if c: parts.append(HUNDS[c])
        if 1 <= r <= 19:
            parts.append(ONES[r])
        elif r == 20:
            parts.append("VEINTE")
        elif 21 <= r <= 29:
            parts.append(VEINTI[r])
        elif r >= 30:
            d, u = divmod(r, 10)
            parts.append(TENS[d] + (f" Y {ONES[u]}" if u else ""))
        return " ".join(parts)

    resultado = []
    if n >= 1_000_000_000_000:
        b = n // 1_000_000_000_000; n %= 1_000_000_000_000
        resultado.append(f"{_grupo(b)} {'BILLÓN' if b == 1 else 'BILLONES'}")
    if n >= 1_000_000_000:
        mm = n // 1_000_000_000; n %= 1_000_000_000
        resultado.append("MIL MILLONES" if mm == 1 else f"{_grupo(mm)} MIL MILLONES")
    if n >= 1_000_000:
        m = n // 1_000_000; n %= 1_000_000
        resultado.append(f"{_grupo(m)} {'MILLÓN' if m == 1 else 'MILLONES'}")
    if n >= 1_000:
        k = n // 1_000; n %= 1_000
        resultado.append("MIL" if k == 1 else f"{_grupo(k)} MIL")
    if n > 0:
        resultado.append(_grupo(n))

    prefijo = "MENOS " if valor < 0 else ""
    return prefijo + " ".join(resultado) + " PESOS M/CTE"


def _fmt_cop(v: float) -> str:
    """Formato colombiano: $1.234.567"""
    return f"${int(round(v)):,}".replace(",", ".")


def _fmt_fecha() -> str:
    hoy = date.today()
    return f"{hoy.day} de {MESES_ES[hoy.month-1]} de {hoy.year}"


def _cargar_perfil(cliente_id: str) -> dict:
    """Carga perfil del cliente: /tmp → Supabase."""
    cache = Path(f"/tmp/siaco/{cliente_id}/perfil.json")
    if cache.exists():
        try:
            return json.loads(cache.read_text(encoding="utf-8"))
        except Exception:
            pass
    try:
        from supabase_client import sb_download
        data = sb_download(f"clientes/{cliente_id}/perfil.json")
        if data:
            perfil = json.loads(data.decode("utf-8"))
            cache.parent.mkdir(parents=True, exist_ok=True)
            cache.write_bytes(data)
            return perfil
    except Exception:
        pass
    return {}


def _cargar_proceso(cliente_id: str, proceso_id: str) -> dict:
    """Busca datos del proceso en expedientes (3-layer persistence)."""
    try:
        from gestor_expedientes import listar_expedientes
        for exp in listar_expedientes(cliente_id):
            if exp.get("proceso_id") == proceso_id:
                return exp
    except Exception:
        pass
    return {}


def _cargar_params_sesion(cliente_id: str) -> dict:
    """Lee parámetros del proceso guardados en sesión por analizar_pliego."""
    try:
        from contexto_sesion import obtener_contexto_sesion
        ctx = obtener_contexto_sesion(cliente_id)
        return ctx.get("parametros_proceso", {})
    except Exception:
        return {}


def _vars_base(cliente_id: str, proceso_id: str) -> dict:
    """Arma diccionario de sustitución con datos del cliente y proceso."""
    perfil  = _cargar_perfil(cliente_id)
    proceso = _cargar_proceso(cliente_id, proceso_id)
    sesion  = _cargar_params_sesion(cliente_id)

    fin  = perfil.get("financiero", {})
    patr = float(fin.get("patrimonio_liquido", perfil.get("patrimonio_liquido", 0)))

    nombre_entidad = (proceso.get("entidad")
                      or sesion.get("entidad", "[ENTIDAD CONTRATANTE]"))
    objeto_proceso = (proceso.get("nombre")
                      or proceso.get("objeto")
                      or sesion.get("objeto", "[OBJETO DEL CONTRATO]"))
    plazo = str(proceso.get("plazo_meses", sesion.get("plazo_meses", "[PLAZO]")))

    return {
        "razon_social":         perfil.get("nombre", perfil.get("razon_social", "[RAZÓN SOCIAL]")),
        "nit":                  perfil.get("nit", "[NIT]"),
        "nombre_representante": perfil.get("representante_legal", "[REPRESENTANTE LEGAL]"),
        "cargo_representante":  perfil.get("cargo_representante", "Representante Legal"),
        "cedula_representante": perfil.get("cedula_representante", "[C.C.]"),
        "direccion":            perfil.get("direccion", "[DIRECCIÓN]"),
        "telefono":             perfil.get("telefono", perfil.get("contacto_whatsapp", "[TELÉFONO]")),
        "correo":               perfil.get("correo", perfil.get("contacto_email", "[CORREO]")),
        "ciudad":               perfil.get("ciudad", "Bogotá D.C."),
        "fecha":                _fmt_fecha(),
        "codigo_proceso":       proceso_id or "[CÓDIGO PROCESO]",
        "nombre_entidad":       nombre_entidad,
        "objeto_proceso":       objeto_proceso,
        "plazo_meses":          plazo,
        "plazo_meses_num":      plazo,
        "patrimonio_liquido":   _fmt_cop(patr),
    }


def _sustituir(plantilla: str, variables: dict) -> str:
    """Reemplaza todos los marcadores {campo} en la plantilla."""
    for k, v in variables.items():
        plantilla = plantilla.replace("{" + k + "}", str(v) if v is not None else "")
    return plantilla


def _generar_docx(titulo: str, contenido: str, empresa: str = "", proceso: str = "") -> bytes:
    """
    Crea .docx profesional: TNR 12pt, interlineado 1.5, texto justificado,
    encabezado empresa|proceso|Pág.X/Y y pie de página con disclaimer SIACO.
    """
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx no disponible")

    doc = Document()

    # Márgenes 2.5 cm todos los lados
    for sec in doc.sections:
        sec.top_margin    = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    # Estilo base: TNR 12pt, justificado, interlineado 1.5, sin espacio extra
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.alignment         = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after       = Pt(0)
    pf.space_before      = Pt(0)

    # Encabezado: empresa | proceso | Pág. X de Y
    if empresa or proceso:
        hdr = doc.sections[0].header
        hp  = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.clear()

        def _hr(text: str):
            r = hp.add_run(text)
            r.font.name      = 'Times New Roman'
            r.font.size      = Pt(9)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        def _field(code: str):
            r = hp.add_run()
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9)
            fc1 = OxmlElement('w:fldChar'); fc1.set(qn('w:fldCharType'), 'begin')
            it  = OxmlElement('w:instrText')
            it.text = f' {code} '
            it.set(qn('xml:space'), 'preserve')
            fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'end')
            r._r.extend([fc1, it, fc2])

        partes = [p[:35] for p in [empresa, proceso] if p]
        _hr("  |  ".join(partes) + "  |  Pág. ")
        _field("PAGE")
        _hr(" de ")
        _field("NUMPAGES")

    # Pie de página
    ftr = doc.sections[0].footer
    fp  = ftr.paragraphs[0] if ftr.paragraphs else ftr.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.clear()
    fr = fp.add_run(
        "Generado con SIACO — Borrador de referencia. "
        "Verifique con asesor jurídico antes de presentar en SECOP II."
    )
    fr.font.name      = 'Times New Roman'
    fr.font.size      = Pt(8)
    fr.font.italic    = True
    fr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    # Título principal
    t = doc.add_heading(titulo, level=1)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in t.runs:
        run.font.name      = 'Times New Roman'
        run.font.size      = Pt(14)
        run.font.bold      = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    doc.add_paragraph()

    # Contenido línea por línea
    for linea in contenido.split("\n"):
        raw      = linea.rstrip()
        stripped = raw.strip()

        if stripped and all(c in "━─═" for c in stripped):
            p = doc.add_paragraph()
            r = p.add_run("─" * 72)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        elif stripped.startswith("⚠"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(stripped.lstrip("⚠️ ").strip())
            r.font.name      = 'Times New Roman'
            r.font.size      = Pt(10)
            r.font.bold      = True
            r.font.color.rgb = RGBColor(0xBF, 0x16, 0x00)

        elif not stripped:
            doc.add_paragraph()

        elif (stripped.isupper()
              and len(stripped) > 5
              and not stripped.startswith("$")
              and not stripped.startswith("COP")
              and not stripped[0].isdigit()):
            # Encabezados de sección en mayúsculas
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            r = p.add_run(raw)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.bold = True

        else:
            p = doc.add_paragraph()
            r = p.add_run(raw)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _tabla_docx_items(items: list) -> str:
    """Construye tabla de ítems APU en texto plano para la plantilla."""
    lineas = []
    for i, it in enumerate(items, 1):
        nom  = it.get("nombre", it.get("descripcion", f"Ítem {i}"))
        und  = it.get("unidad", "und")
        cant = it.get("cantidad", 1)
        pu   = it.get("precio_unitario", 0)
        tot  = cant * pu
        lineas.append(
            f"  {i:>3}. {nom:<35} {und:<6} {cant:>8.2f}  {_fmt_cop(pu):>16}  {_fmt_cop(tot):>16}"
        )
    return "\n".join(lineas) if lineas else "  [Sin ítems — ingrese desde la Calculadora APU]"


# ── Modelos ──────────────────────────────────────────────────────────────────

class DocRequest(BaseModel):
    cliente_id: str
    proceso_id: str

class FormularioEconomicoRequest(BaseModel):
    cliente_id: str
    proceso_id: str
    datos_apu:  dict = {}

class CapacidadResidualRequest(BaseModel):
    cliente_id:         str
    proceso_id:         str
    k_requerido:        float = 0.0
    contratos_vigentes: list[dict] = []

class PaqueteRequest(BaseModel):
    cliente_id:         str
    proceso_id:         str
    datos_apu:          dict = {}
    k_requerido:        float = 0.0
    contratos_vigentes: list[dict] = []


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.post("/oferta/carta-presentacion")
def carta_presentacion(body: DocRequest, authorization: str = Header(None)):
    """Genera Folio 1 — Carta de presentación con 6 declaraciones juramentadas en Word."""
    from routers.auth import require_auth
    require_auth(authorization)

    vars_ = _vars_base(body.cliente_id, body.proceso_id)
    vars_["valor_letras"]  = "[VALOR EN LETRAS] PESOS M/CTE"
    vars_["valor_numeros"] = "$ [VALOR OFERTA]"

    tpl        = (FORMULARIOS_DIR / "carta_presentacion.txt").read_text(encoding="utf-8")
    contenido  = _sustituir(tpl, vars_)
    docx_bytes = _generar_docx(
        "CARTA DE PRESENTACIÓN DE OFERTA",
        contenido,
        empresa=vars_.get("razon_social", ""),
        proceso=body.proceso_id,
    )
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":
                 f'attachment; filename="01_Carta_Presentacion_{body.proceso_id}.docx"'},
    )


@router.post("/oferta/formulario-economico")
def formulario_economico(body: FormularioEconomicoRequest, authorization: str = Header(None)):
    """Genera Folio 7 — Propuesta Económica con estructura AIU-CCE en Word."""
    from routers.auth import require_auth
    require_auth(authorization)

    vars_ = _vars_base(body.cliente_id, body.proceso_id)
    apu   = body.datos_apu
    cd    = apu.get("costos_directos", {})
    aiu_d = apu.get("aiu", {})

    materiales  = apu.get("items_materiales",  apu.get("materiales_lista", []))
    personal    = apu.get("items_personal",    apu.get("personal_lista",   []))
    equipos     = apu.get("items_equipos",     apu.get("equipos_lista",    []))
    todos_items = list(materiales) + list(personal) + list(equipos)

    precio_total = float(apu.get("precio_minimo", 0))

    vars_.update({
        "tabla_items":            _tabla_docx_items(todos_items),
        "subtotal_directos":      _fmt_cop(float(cd.get("subtotal", 0))),
        "admin_pct":              f"{aiu_d.get('admin_pct', 12):.1f}",
        "valor_admin":            _fmt_cop(float(aiu_d.get("administracion", 0))),
        "imprev_pct":             f"{aiu_d.get('imprevistos_pct', 3):.1f}",
        "valor_imprev":           _fmt_cop(float(aiu_d.get("imprevistos", 0))),
        "util_pct":               f"{aiu_d.get('utilidad_pct', 8):.1f}",
        "valor_util":             _fmt_cop(float(aiu_d.get("utilidad", 0))),
        "subtotal_aiu":           _fmt_cop(float(aiu_d.get("total_aiu", 0))),
        "valor_total":            _fmt_cop(precio_total),
        "valor_iva":              _fmt_cop(0),
        "valor_total_definitivo": _fmt_cop(precio_total),
        "valor_letras":           numero_a_letras(precio_total) if precio_total else "[VALOR EN LETRAS]",
    })

    tpl        = (FORMULARIOS_DIR / "formulario_economico.txt").read_text(encoding="utf-8")
    contenido  = _sustituir(tpl, vars_)
    docx_bytes = _generar_docx(
        "FORMULARIO 1 – PROPUESTA ECONÓMICA",
        contenido,
        empresa=vars_.get("razon_social", ""),
        proceso=body.proceso_id,
    )
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":
                 f'attachment; filename="02_Formulario_Economico_{body.proceso_id}.docx"'},
    )


@router.post("/oferta/formato-experiencia")
def formato_experiencia(body: DocRequest, authorization: str = Header(None)):
    """Genera Folio 3 — Experiencia del Proponente con UNSPSC y Cert. No. en Word."""
    from routers.auth import require_auth
    require_auth(authorization)

    vars_     = _vars_base(body.cliente_id, body.proceso_id)
    perfil    = _cargar_perfil(body.cliente_id)
    contratos = perfil.get("contratos_experiencia", perfil.get("experiencia_contratos", []))
    exp_agg   = perfil.get("experiencia", {})

    if isinstance(contratos, list) and contratos:
        bloques = []
        for i, c in enumerate(contratos, 1):
            bloques.append(
                f"  CONTRATO {i}\n"
                f"    Entidad/Cliente   : {c.get('entidad', c.get('cliente', '[INSERTAR_ENTIDAD]'))}\n"
                f"    Objeto            : {c.get('objeto', '[INSERTAR_OBJETO]')}\n"
                f"    Código UNSPSC     : {c.get('unspsc', c.get('codigo_unspsc', '[INSERTAR_UNSPSC]'))}\n"
                f"    Valor total       : {_fmt_cop(float(c.get('valor', 0)))}\n"
                f"    Período           : {c.get('fecha_inicio', c.get('inicio', '[INICIO]'))} "
                f"— {c.get('fecha_fin', c.get('fin', '[FIN]'))}\n"
                f"    % Participación   : {c.get('porcentaje_participacion', c.get('participacion', 100))}%\n"
                f"    Certificado No.   : {c.get('certificado_no', '[INSERTAR_No_CERTIFICADO]')}"
            )
        tabla_exp = "\n\n".join(bloques)

    elif isinstance(exp_agg, dict) and exp_agg.get("valor_acumulado"):
        tabla_exp = (
            "  DATOS DEL PERFIL SIACO (complete un bloque por contrato):\n\n"
            "  CONTRATO 1\n"
            f"    Entidad/Cliente   : [INSERTAR_ENTIDAD]\n"
            f"    Objeto            : {exp_agg.get('objeto_similar', '[INSERTAR_OBJETO]')}\n"
            f"    Código UNSPSC     : {exp_agg.get('codigos_unspsc', '[INSERTAR_UNSPSC]')}\n"
            f"    Valor total       : {_fmt_cop(float(exp_agg.get('valor_individual_max', 0)))} "
            "(mayor contrato individual)\n"
            "    Período           : [INSERTAR_FECHAS]\n"
            "    % Participación   : [INSERTAR_%]\n"
            "    Certificado No.   : [INSERTAR_No_CERTIFICADO]\n\n"
            "  ⚠️ Complete un bloque por cada contrato. "
            "Adjunte los certificados como soporte."
        )
    else:
        tabla_exp = (
            "  CONTRATO 1\n"
            "    Entidad/Cliente   : [INSERTAR_ENTIDAD]\n"
            "    Objeto            : [INSERTAR_OBJETO]\n"
            "    Código UNSPSC     : [INSERTAR_UNSPSC]\n"
            "    Valor total       : $[INSERTAR_VALOR]\n"
            "    Período           : [INSERTAR_FECHA_INICIO] — [INSERTAR_FECHA_FIN]\n"
            "    % Participación   : [INSERTAR_%]\n"
            "    Certificado No.   : [INSERTAR_No_CERTIFICADO]\n\n"
            "  Complete un bloque por cada contrato que acredita experiencia.\n"
            "  Adjunte certificados originales de cada entidad contratante."
        )

    vars_["tabla_experiencia"] = tabla_exp
    tpl        = (FORMULARIOS_DIR / "formulario_experiencia.txt").read_text(encoding="utf-8")
    contenido  = _sustituir(tpl, vars_)
    docx_bytes = _generar_docx(
        "FORMULARIO 2 – EXPERIENCIA DEL PROPONENTE",
        contenido,
        empresa=vars_.get("razon_social", ""),
        proceso=body.proceso_id,
    )
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":
                 f'attachment; filename="03_Formulario_Experiencia_{body.proceso_id}.docx"'},
    )


@router.post("/oferta/capacidad-residual")
def capacidad_residual(body: CapacidadResidualRequest, authorization: str = Header(None)):
    """Genera Folio 4 — Factor K: Capacidad Residual de Contratación en Word."""
    from routers.auth import require_auth
    require_auth(authorization)

    vars_  = _vars_base(body.cliente_id, body.proceso_id)
    perfil = _cargar_perfil(body.cliente_id)
    fin    = perfil.get("financiero", {})
    patr   = float(fin.get("patrimonio_liquido", perfil.get("patrimonio_liquido", 0)))
    ochenta = patr * 0.8

    vigentes   = body.contratos_vigentes or perfil.get("contratos_en_ejecucion", [])
    total_comp = sum(
        float(c.get("valor", 0)) * float(c.get("pct_pendiente", c.get("porcentaje_pendiente", 0))) / 100
        for c in vigentes
    )
    k_residual = ochenta - total_comp
    k_req      = body.k_requerido

    if vigentes:
        filas = ["  Contrato / Objeto                    | Valor total    | % pendiente | Compromiso"]
        filas.append("  " + "-" * 78)
        for c in vigentes:
            obj  = str(c.get("objeto", c.get("descripcion", "")))[:38]
            val  = float(c.get("valor", 0))
            pct  = float(c.get("pct_pendiente", c.get("porcentaje_pendiente", 0)))
            comp = val * pct / 100
            filas.append(
                f"  {obj:<38} | {_fmt_cop(val):>14} | {pct:>11.1f}% | {_fmt_cop(comp):>14}"
            )
        tabla_v = "\n".join(filas)
    else:
        tabla_v = "  [Sin contratos en ejecución — ingrese manualmente si aplica]"

    resultado_k = ("CUMPLE" if k_residual >= k_req
                   else "NO CUMPLE — Verifique contratos en ejecucion y patrimonio en RUP")

    vars_.update({
        "tabla_contratos_ejecucion": tabla_v,
        "total_compromisos":         _fmt_cop(total_comp),
        "ochenta_pct_patrimonio":    _fmt_cop(ochenta),
        "k_residual":                _fmt_cop(k_residual),
        "k_requerido":               _fmt_cop(k_req) if k_req else "[Ver pliego]",
        "resultado_k":               resultado_k,
    })

    tpl        = (FORMULARIOS_DIR / "formato_k_residual.txt").read_text(encoding="utf-8")
    contenido  = _sustituir(tpl, vars_)
    docx_bytes = _generar_docx(
        "CAPACIDAD RESIDUAL DE CONTRATACIÓN — FACTOR K",
        contenido,
        empresa=vars_.get("razon_social", ""),
        proceso=body.proceso_id,
    )
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":
                 f'attachment; filename="04_Capacidad_Residual_{body.proceso_id}.docx"'},
    )


@router.post("/oferta/paz-y-salvos")
def paz_y_salvos(body: DocRequest, authorization: str = Header(None)):
    """Genera Folio 8 — Instrucciones para obtención de paz y salvos parafiscales."""
    from routers.auth import require_auth
    require_auth(authorization)

    vars_      = _vars_base(body.cliente_id, body.proceso_id)
    tpl        = (FORMULARIOS_DIR / "paz_y_salvos.txt").read_text(encoding="utf-8")
    contenido  = _sustituir(tpl, vars_)
    docx_bytes = _generar_docx(
        "INSTRUCCIONES PAZ Y SALVOS PARAFISCALES",
        contenido,
        empresa=vars_.get("razon_social", ""),
        proceso=body.proceso_id,
    )
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":
                 f'attachment; filename="05_Paz_Y_Salvos_{body.proceso_id}.docx"'},
    )


@router.post("/oferta/checklist")
def checklist_oferta(body: DocRequest, authorization: str = Header(None)):
    """Genera checklist anti-rechazo final de la oferta."""
    from routers.auth import require_auth
    require_auth(authorization)

    vars_      = _vars_base(body.cliente_id, body.proceso_id)
    tpl        = (FORMULARIOS_DIR / "checklist_oferta.txt").read_text(encoding="utf-8")
    contenido  = _sustituir(tpl, vars_)
    docx_bytes = _generar_docx(
        "CHECKLIST ANTI-RECHAZO — PAQUETE DE OFERTA",
        contenido,
        empresa=vars_.get("razon_social", ""),
        proceso=body.proceso_id,
    )
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":
                 f'attachment; filename="00_Checklist_Anti_Rechazo_{body.proceso_id}.docx"'},
    )


@router.post("/oferta/paquete-completo")
def paquete_completo(body: PaqueteRequest, authorization: str = Header(None)):
    """Genera ZIP con estructura de carpetas DOCUMENTOS_SIACO / DOCUMENTOS_USUARIO."""
    from routers.auth import require_auth
    require_auth(authorization)

    def _get_bytes(endpoint_fn, req_body):
        return endpoint_fn(req_body, authorization=authorization).body

    vars_ = _vars_base(body.cliente_id, body.proceso_id)
    emp   = vars_.get("razon_social", "empresa")
    pid   = body.proceso_id
    siaco = f"OFERTA_{pid}/DOCUMENTOS_SIACO/"

    docs: dict[str, bytes] = {}

    docs[f"{siaco}01_Carta_Presentacion.docx"] = _get_bytes(
        carta_presentacion, DocRequest(cliente_id=body.cliente_id, proceso_id=pid))

    docs[f"{siaco}02_Formulario_Economico.docx"] = _get_bytes(
        formulario_economico, FormularioEconomicoRequest(
            cliente_id=body.cliente_id, proceso_id=pid, datos_apu=body.datos_apu))

    docs[f"{siaco}03_Formulario_Experiencia.docx"] = _get_bytes(
        formato_experiencia, DocRequest(cliente_id=body.cliente_id, proceso_id=pid))

    docs[f"{siaco}04_Capacidad_Residual.docx"] = _get_bytes(
        capacidad_residual, CapacidadResidualRequest(
            cliente_id=body.cliente_id, proceso_id=pid,
            k_requerido=body.k_requerido, contratos_vigentes=body.contratos_vigentes))

    docs[f"{siaco}05_Paz_Y_Salvos_Instrucciones.docx"] = _get_bytes(
        paz_y_salvos, DocRequest(cliente_id=body.cliente_id, proceso_id=pid))

    docs[f"{siaco}06_Checklist_Anti_Rechazo.docx"] = _get_bytes(
        checklist_oferta, DocRequest(cliente_id=body.cliente_id, proceso_id=pid))

    # README para carpeta de docs del usuario
    readme = (
        "DOCUMENTOS QUE USTED DEBE APORTAR\n"
        "====================================\n\n"
        "Coloque aquí los documentos habilitantes escaneados en PDF:\n\n"
        "  06_RUT.pdf\n"
        "  07_Camara_Comercio.pdf         (máx 30 días de expedición)\n"
        "  08_Cedula_Representante.pdf\n"
        "  09_Antecedentes_Disciplinarios.pdf\n"
        "  10_Antecedentes_Fiscales.pdf\n"
        "  11_Antecedentes_Penales.pdf\n"
        "  12_Paz_Salvo_SENA.pdf\n"
        "  13_Paz_Salvo_ICBF.pdf\n"
        "  14_Paz_Salvo_Caja.pdf\n"
        "  15_Planilla_PILA.pdf\n"
        "  16_Garantia_Seriedad.pdf\n"
        "  17_Certificados_Experiencia.pdf\n\n"
        "Instrucciones: DOCUMENTOS_SIACO/05_Paz_Y_Salvos_Instrucciones.docx\n"
        "Checklist    : DOCUMENTOS_SIACO/06_Checklist_Anti_Rechazo.docx\n"
    )
    docs[f"OFERTA_{pid}/DOCUMENTOS_USUARIO/README.txt"] = readme.encode("utf-8")

    # Índice como Word
    indice_txt = (
        f"Proceso: {pid}\n"
        f"Empresa: {emp}\n"
        f"Generado: {_fmt_fecha()}\n\n"
        f"DOCUMENTOS SIACO (generados automáticamente):\n"
        f"  01_Carta_Presentacion.docx    — Carta con 6 declaraciones juramentadas\n"
        f"  02_Formulario_Economico.docx  — Propuesta económica con AIU-CCE\n"
        f"  03_Formulario_Experiencia.docx— Experiencia con UNSPSC y Cert. No.\n"
        f"  04_Capacidad_Residual.docx    — Factor K residual de contratación\n"
        f"  05_Paz_Y_Salvos_Instrucciones.docx — Guía para obtener paz y salvos\n"
        f"  06_Checklist_Anti_Rechazo.docx    — Verificación final anti-rechazo\n\n"
        f"DOCUMENTOS USUARIO (a su cargo):\n"
        f"  06_RUT | 07_Camara_Comercio | 08_Cedula_Representante\n"
        f"  09_Antecedentes_Disciplinarios | 10_Antecedentes_Fiscales\n"
        f"  11_Antecedentes_Penales | 12-14_Paz_Salvos | 15_PILA\n"
        f"  16_Garantia_Seriedad | 17_Certificados_Experiencia\n\n"
        f"ADVERTENCIA: Borradores generados automáticamente por SIACO.\n"
        f"Revise, complete y firme cada documento antes de presentar en SECOP II."
    )
    docs[f"OFERTA_{pid}/00_INDICE_PAQUETE.docx"] = _generar_docx(
        f"ÍNDICE — PAQUETE DE OFERTA\n{pid}",
        indice_txt, empresa=emp, proceso=pid,
    )

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for nombre, contenido in docs.items():
            zf.writestr(nombre, contenido)
    buf.seek(0)

    fecha_str = date.today().isoformat()
    return Response(
        content=buf.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition":
                 f'attachment; filename="SIACO_Oferta_{pid}_{fecha_str}.zip"'},
    )
