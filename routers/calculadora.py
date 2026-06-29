# -*- coding: utf-8 -*-
"""Router de Calculadora de Oferta Económica — SIACO v3.0"""
import io
from datetime import date
from fastapi import APIRouter, HTTPException, Header
from fastapi.responses import Response
from pydantic import BaseModel, Field
from typing import Optional

from calculadora_apu import CalculadoraAPU

router = APIRouter(tags=["calculadora"])
_calc = CalculadoraAPU()


# ── Modelos de entrada ────────────────────────────────────────────────────────

class TrabajadorBody(BaseModel):
    salario_basico: float = Field(..., gt=0)
    nivel_riesgo:   int   = Field(4, ge=1, le=5)


class MaterialItem(BaseModel):
    nombre:          str   = ""
    unidad:          str   = "und"
    cantidad:        float = Field(1.0, ge=0)
    precio_unitario: float = Field(0.0, ge=0)
    iva:             bool  = False


class PersonalItem(BaseModel):
    cargo:        str   = ""
    salario:      float = Field(0.0, ge=0)
    cantidad:     int   = Field(1, ge=1)
    dias:         float = Field(30.0, ge=1)
    nivel_riesgo: int   = Field(4, ge=1, le=5)


class EquipoItem(BaseModel):
    nombre:    str   = ""
    costo_dia: float = Field(0.0, ge=0)
    dias:      float = Field(1.0, ge=1)


class AIUParams(BaseModel):
    admin_pct:       float = Field(0.12, ge=0, le=1)
    imprevistos_pct: float = Field(0.03, ge=0, le=1)
    utilidad_pct:    float = Field(0.08, ge=0, le=1)


class OfertaBody(BaseModel):
    materiales:          list[MaterialItem] = []
    personal:            list[PersonalItem] = []
    equipos:             list[EquipoItem]   = []
    aiu:                 AIUParams          = AIUParams()
    plazo_meses:         int                = Field(6, ge=1, le=60)
    incluir_polizas:     bool               = True
    presupuesto_oficial: Optional[float]    = None


class FlujoCajaBody(BaseModel):
    valor_contrato:  float = Field(..., gt=0)
    anticipo_pct:    float = Field(0.30, ge=0, le=0.50)
    plazo_meses:     int   = Field(6, ge=1, le=60)
    costos_directos: float = Field(0.0, ge=0)
    aiu_pct:         float = Field(0.23, ge=0, le=1)


class ExportarExcelBody(BaseModel):
    materiales: list[dict] = []
    personal:   list[dict] = []
    equipos:    list[dict] = []
    resultado:  dict       = {}
    flujo:      dict       = {}


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.post("/calculadora/trabajador")
def costo_trabajador(body: TrabajadorBody, authorization: str = Header(None)):
    """Costo real mensual de un trabajador con desglose de prestaciones sociales."""
    from routers.auth import require_auth
    require_auth(authorization)
    try:
        return _calc.calcular_costo_real_trabajador(body.salario_basico, body.nivel_riesgo)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/calculadora/hora-apu")
def hora_apu(body: TrabajadorBody, authorization: str = Header(None)):
    """Valor hora-hombre APU (metodología IDU/INVIAS) para presupuestos de construcción."""
    from routers.auth import require_auth
    require_auth(authorization)
    try:
        return _calc.calcular_valor_hora_apu(body.salario_basico, body.nivel_riesgo)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/calculadora/oferta")
def calcular_oferta(body: OfertaBody, authorization: str = Header(None)):
    """
    Precio mínimo y sugerido de una oferta económica con desglose completo.
    Si se proporciona presupuesto_oficial incluye análisis de viabilidad.
    """
    from routers.auth import require_auth
    require_auth(authorization)
    try:
        resultado = _calc.calcular_precio_oferta(
            materiales      = [m.model_dump() for m in body.materiales],
            personal        = [p.model_dump() for p in body.personal],
            equipos         = [e.model_dump() for e in body.equipos],
            admin_pct       = body.aiu.admin_pct,
            imprevistos_pct = body.aiu.imprevistos_pct,
            utilidad_pct    = body.aiu.utilidad_pct,
            incluir_polizas = body.incluir_polizas,
            plazo_meses     = body.plazo_meses,
        )
        if body.presupuesto_oficial:
            resultado["viabilidad"] = _calc.verificar_viabilidad_presupuesto(
                precio_minimo       = resultado["precio_minimo"],
                presupuesto_oficial = body.presupuesto_oficial,
            )
        return resultado
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/calculadora/flujo-caja")
def flujo_caja(body: FlujoCajaBody, authorization: str = Header(None)):
    """Proyección de flujo de caja según mecánica real de contratos públicos colombianos."""
    from routers.auth import require_auth
    require_auth(authorization)
    try:
        return _calc.calcular_flujo_caja(
            valor_contrato  = body.valor_contrato,
            anticipo_pct    = body.anticipo_pct,
            plazo_meses     = body.plazo_meses,
            costos_directos = body.costos_directos,
            aiu_pct         = body.aiu_pct,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/calculadora/exportar-excel")
def exportar_excel(body: ExportarExcelBody, authorization: str = Header(None)):
    """Genera archivo .xlsx profesional multi-hoja con la oferta económica completa."""
    from routers.auth import require_auth
    require_auth(authorization)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.chart import PieChart, BarChart, Reference
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no disponible")

    try:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # eliminar hoja por defecto

        # ── Paleta corporativa ────────────────────────────────────────────
        C_AZUL        = "1F4E79"
        C_AZUL_MED    = "2E75B6"
        C_AZUL_CLR    = "DEEAF1"
        C_AZUL_ALT    = "BDD7EE"
        C_VERDE_CLR   = "E2EFDA"
        C_AMARILLO    = "FFF2CC"
        C_ROJO_CLR    = "FFE0E0"
        C_BLANCO      = "FFFFFF"
        C_GRIS_CLR    = "F2F2F2"
        FMT_PESOS     = "$#,##0"
        FMT_PCT       = "0.00%"
        FMT_NUM       = "#,##0.00"

        def _fnt(bold=False, size=11, color="000000", italic=False):
            return Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)

        def _fil(color):
            return PatternFill("solid", fgColor=color)

        def _brd(thin=True):
            s = Side(style="thin" if thin else None)
            return Border(left=s, right=s, top=s, bottom=s)

        def _aln(h="left", v="center", wrap=False):
            return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

        def _hdr_cells(ws, values: list, row_n: int):
            """Aplica estilo de encabezado a una fila completa."""
            for col, val in enumerate(values, 1):
                c = ws.cell(row=row_n, column=col, value=val)
                c.font      = _fnt(bold=True, color=C_BLANCO)
                c.fill      = _fil(C_AZUL)
                c.border    = _brd()
                c.alignment = _aln("center")

        def _data_cell(ws, row_n, col_n, value, fmt=None, bold=False, fill=None,
                       align="left", color="000000"):
            c = ws.cell(row=row_n, column=col_n, value=value)
            c.font      = _fnt(bold=bold, color=color)
            c.border    = _brd()
            c.alignment = _aln(align)
            if fmt:  c.number_format = fmt
            if fill: c.fill = _fil(fill)
            return c

        def _auto_width(ws, min_w=10, max_w=45):
            for col in ws.columns:
                w = max((len(str(c.value or "")) for c in col), default=min_w)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 3, max_w)

        def _freeze(ws):
            ws.freeze_panes = "A2"

        res     = body.resultado
        cd      = res.get("costos_directos", {})
        aiu     = res.get("aiu", {})
        rsm     = res.get("resumen", {})
        hoy     = date.today().strftime("%d/%m/%Y")
        n_fname = f"SIACO_Analisis_Oferta_{date.today().isoformat()}.xlsx"

        # ══════════════════════════════════════════════════════════════════
        # HOJA 1 — Resumen Ejecutivo
        # ══════════════════════════════════════════════════════════════════
        ws1 = wb.create_sheet("Resumen Ejecutivo")

        # Logo / título
        c = ws1.cell(row=1, column=1, value="SIACO")
        c.font = _fnt(bold=True, size=24, color=C_AZUL)
        ws1.merge_cells("A1:F1")

        c = ws1.cell(row=2, column=1, value="ANÁLISIS DE PRECIO DE OFERTA — Metodología APU 2026")
        c.font = _fnt(bold=True, size=13, color=C_AZUL_MED)
        ws1.merge_cells("A2:F2")

        ws1.cell(row=3, column=1, value=f"Fecha: {hoy}").font = _fnt(italic=True, color="555555")
        ws1.row_dimensions[4].height = 8

        # Tabla resumen
        hdrs = ["Componente", "Valor (COP)"]
        _hdr_cells(ws1, hdrs, 5)

        rows_res = [
            ("Materiales e insumos",                      cd.get("materiales", 0),          C_BLANCO),
            ("Mano de obra (con prestaciones)",           cd.get("mano_obra", 0),            C_GRIS_CLR),
            ("Equipos y herramientas",                    cd.get("equipos", 0),              C_BLANCO),
            ("Subtotal costos directos",                  cd.get("subtotal", 0),             C_AZUL_CLR),
            (f"Administración ({aiu.get('admin_pct',0)}%)",      aiu.get("administracion", 0), C_BLANCO),
            (f"Imprevistos ({aiu.get('imprevistos_pct',0)}%)",   aiu.get("imprevistos", 0),    C_GRIS_CLR),
            (f"Utilidad ({aiu.get('utilidad_pct',0)}%)",         aiu.get("utilidad", 0),       C_BLANCO),
            ("Subtotal AIU",                              aiu.get("total_aiu", 0),           C_AZUL_ALT),
            ("Pólizas estimadas",                         res.get("polizas_estimadas", 0),   C_GRIS_CLR),
        ]
        for i, (label, valor, fill) in enumerate(rows_res, 6):
            _data_cell(ws1, i, 1, label,  fill=fill)
            _data_cell(ws1, i, 2, valor,  fmt=FMT_PESOS, align="right", fill=fill)

        row_sep = 6 + len(rows_res)
        ws1.row_dimensions[row_sep].height = 8

        # Destacados: precio mínimo / sugerido / utilidad
        destacados = [
            ("PRECIO MÍNIMO (no perder dinero)",   res.get("precio_minimo", 0),              C_AMARILLO, "FFC000"),
            ("PRECIO SUGERIDO (+5 % de seguridad)",res.get("precio_sugerido", 0),            C_VERDE_CLR, "70AD47"),
            ("UTILIDAD PROYECTADA",                rsm.get("utilidad_proyectada", 0),        C_AZUL_CLR,  C_AZUL_MED),
            ("MARGEN DE UTILIDAD",                 f"{rsm.get('margen_utilidad_pct', 0)} %", C_GRIS_CLR, "555555"),
        ]
        for j, (label, valor, fill, col_txt) in enumerate(destacados, row_sep + 1):
            c1 = ws1.cell(row=j, column=1, value=label)
            c1.font      = _fnt(bold=True, size=12, color=col_txt)
            c1.fill      = _fil(fill)
            c1.border    = Border(
                left=Side(style="medium", color=col_txt),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )
            c1.alignment = _aln("left")
            c2 = ws1.cell(row=j, column=2, value=valor)
            c2.font      = _fnt(bold=True, size=12, color=col_txt)
            c2.fill      = _fil(fill)
            c2.border    = _brd()
            c2.alignment = _aln("right")
            if isinstance(valor, (int, float)):
                c2.number_format = FMT_PESOS

        ws1.column_dimensions["A"].width = 42
        ws1.column_dimensions["B"].width = 20
        _freeze(ws1)

        # ══════════════════════════════════════════════════════════════════
        # HOJA 2 — Materiales
        # ══════════════════════════════════════════════════════════════════
        ws2 = wb.create_sheet("Materiales")
        hdrs2 = ["#", "Descripción", "Unidad", "Cantidad", "Precio Unitario", "Total"]
        _hdr_cells(ws2, hdrs2, 1)

        total_mat = 0
        for i, m in enumerate(body.materiales, 1):
            sub  = (m.get("cantidad", 0) or 0) * (m.get("precio_unitario", 0) or 0)
            fill = C_BLANCO if i % 2 == 0 else C_GRIS_CLR
            _data_cell(ws2, i+1, 1, i,                            fill=fill, align="center")
            _data_cell(ws2, i+1, 2, m.get("nombre",""),           fill=fill)
            _data_cell(ws2, i+1, 3, m.get("unidad","und"),        fill=fill, align="center")
            _data_cell(ws2, i+1, 4, m.get("cantidad",0),          fill=fill, align="right", fmt=FMT_NUM)
            _data_cell(ws2, i+1, 5, m.get("precio_unitario",0),   fill=fill, align="right", fmt=FMT_PESOS)
            _data_cell(ws2, i+1, 6, sub,                          fill=fill, align="right", fmt=FMT_PESOS)
            total_mat += sub

        fila_tot = len(body.materiales) + 2
        ws2.merge_cells(f"A{fila_tot}:E{fila_tot}")
        _data_cell(ws2, fila_tot, 1, "TOTAL MATERIALES", bold=True, fill=C_AMARILLO, align="right")
        _data_cell(ws2, fila_tot, 6, total_mat, fmt=FMT_PESOS, bold=True, fill=C_AMARILLO, align="right")

        ws2.column_dimensions["B"].width = 38
        ws2.column_dimensions["C"].width = 10
        ws2.column_dimensions["D"].width = 12
        ws2.column_dimensions["E"].width = 18
        ws2.column_dimensions["F"].width = 18
        _freeze(ws2)

        # ══════════════════════════════════════════════════════════════════
        # HOJA 3 — Mano de Obra
        # ══════════════════════════════════════════════════════════════════
        ws3 = wb.create_sheet("Mano de Obra")
        hdrs3 = ["Cargo", "Salario Básico", "Cant.", "Días", "Factor Prest.", "V/Hora APU", "Costo Total"]
        _hdr_cells(ws3, hdrs3, 1)

        total_mo = 0
        for i, p in enumerate(body.personal, 1):
            sal  = float(p.get("salario", 0))
            cant = int(p.get("cantidad", 1))
            dias = float(p.get("dias", 30))
            cmes = _calc.calcular_costo_real_trabajador(sal)["costo_total_mes"]
            hora = _calc.calcular_valor_hora_apu(sal)["valor_hora_apu"]
            tot  = cmes * cant * (dias / 30)
            fp   = ((cmes / sal) - 1) if sal else 0
            fill = C_BLANCO if i % 2 == 0 else C_GRIS_CLR
            _data_cell(ws3, i+1, 1, p.get("cargo",""),  fill=fill)
            _data_cell(ws3, i+1, 2, sal,                fill=fill, align="right", fmt=FMT_PESOS)
            _data_cell(ws3, i+1, 3, cant,               fill=fill, align="center")
            _data_cell(ws3, i+1, 4, dias,               fill=fill, align="center")
            _data_cell(ws3, i+1, 5, fp,                 fill=fill, align="right", fmt=FMT_PCT)
            _data_cell(ws3, i+1, 6, hora,               fill=fill, align="right", fmt=FMT_PESOS)
            _data_cell(ws3, i+1, 7, tot,                fill=fill, align="right", fmt=FMT_PESOS)
            total_mo += tot

        fila_tot3 = len(body.personal) + 2
        ws3.merge_cells(f"A{fila_tot3}:F{fila_tot3}")
        _data_cell(ws3, fila_tot3, 1, "TOTAL MANO DE OBRA", bold=True, fill=C_AMARILLO, align="right")
        _data_cell(ws3, fila_tot3, 7, total_mo, fmt=FMT_PESOS, bold=True, fill=C_AMARILLO, align="right")

        # Nota metodológica
        fila_nota = fila_tot3 + 2
        nota = ws3.cell(row=fila_nota, column=1,
            value="Nota: Factor prestacional incluye cesantías, prima, vacaciones, "
                  "salud empleador, pensión empleador, ARL nivel 4, SENA, ICBF y caja compensación. "
                  "Valor hora APU incluye TPNL 22.5 % y Mayor Valor Prestacional 14.72 % (IDU/INVIAS 2026).")
        nota.font      = _fnt(italic=True, size=9, color="555555")
        nota.alignment = Alignment(wrap_text=True)
        ws3.merge_cells(f"A{fila_nota}:G{fila_nota}")
        ws3.row_dimensions[fila_nota].height = 36

        for col, w in zip("ABCDEFG", [30, 16, 8, 8, 14, 14, 16]):
            ws3.column_dimensions[col].width = w
        _freeze(ws3)

        # ══════════════════════════════════════════════════════════════════
        # HOJA 4 — AIU con gráfico de torta
        # ══════════════════════════════════════════════════════════════════
        ws4 = wb.create_sheet("AIU")
        _hdr_cells(ws4, ["Componente", "%", "Valor (COP)"], 1)

        aiu_rows = [
            ("Costos directos",  1.0,                                           cd.get("subtotal", 0)),
            ("Administración",   aiu.get("admin_pct", 0) / 100,                 aiu.get("administracion", 0)),
            ("Imprevistos",      aiu.get("imprevistos_pct", 0) / 100,           aiu.get("imprevistos", 0)),
            ("Utilidad",         aiu.get("utilidad_pct", 0) / 100,              aiu.get("utilidad", 0)),
            ("Pólizas estimadas",0,                                              res.get("polizas_estimadas", 0)),
        ]
        fills_aiu = [C_AZUL_CLR, C_GRIS_CLR, C_BLANCO, C_VERDE_CLR, C_AMARILLO]
        for i, (lbl, pct, val) in enumerate(aiu_rows, 2):
            fill = fills_aiu[i - 2]
            _data_cell(ws4, i, 1, lbl,  fill=fill)
            _data_cell(ws4, i, 2, pct,  fmt=FMT_PCT, fill=fill, align="center")
            _data_cell(ws4, i, 3, val,  fmt=FMT_PESOS, fill=fill, align="right")

        # Fila total
        fila_aiu_tot = len(aiu_rows) + 2
        _data_cell(ws4, fila_aiu_tot, 1, "PRECIO MÍNIMO TOTAL", bold=True, fill=C_AMARILLO)
        _data_cell(ws4, fila_aiu_tot, 3, res.get("precio_minimo", 0),
                   fmt=FMT_PESOS, bold=True, fill=C_AMARILLO, align="right")

        # Gráfico de torta: Costos directos vs AIU
        pie = PieChart()
        pie.title  = "Estructura del precio (Costos vs AIU)"
        pie.style  = 10
        pie.width  = 15
        pie.height = 10
        labels_ref = Reference(ws4, min_col=1, min_row=2, max_row=5)
        data_ref   = Reference(ws4, min_col=3, min_row=2, max_row=5)
        pie.add_data(data_ref)
        pie.set_categories(labels_ref)
        ws4.add_chart(pie, "E2")

        for col, w in zip("ABC", [28, 10, 18]):
            ws4.column_dimensions[col].width = w
        _freeze(ws4)

        # ══════════════════════════════════════════════════════════════════
        # HOJA 5 — Flujo de Caja
        # ══════════════════════════════════════════════════════════════════
        ws5 = wb.create_sheet("Flujo de Caja")
        flujo_data = body.flujo.get("flujo_mensual", body.flujo.get("proyeccion", []))

        if flujo_data:
            # Encabezado resumen
            resumen_flujo = body.flujo.get("resumen", "")
            if resumen_flujo:
                c = ws5.cell(row=1, column=1, value=resumen_flujo)
                c.font      = _fnt(italic=True, size=10)
                c.alignment = Alignment(wrap_text=True)
                ws5.merge_cells("A1:H1")
                ws5.row_dimensions[1].height = 36

            hdrs5 = ["Mes", "Acta bruta", "Amort. anticipo", "Ingreso neto",
                     "Costos directos", "Flujo neto mes", "Saldo acumulado", "Estado"]
            _hdr_cells(ws5, hdrs5, 2)

            for i, f in enumerate(flujo_data, 3):
                ingresos = f.get("ingresos", {})
                egresos  = f.get("egresos", {})
                acta     = ingresos.get("acta_cobrada", ingresos if isinstance(ingresos, (int, float)) else 0)
                amort    = f.get("amortizacion_anticipo", 0)
                neto_ing = acta - amort if isinstance(acta, (int, float)) else 0
                costos   = egresos.get("total_egresos", egresos if isinstance(egresos, (int, float)) else 0)
                flujo_n  = f.get("flujo_neto_mes", f.get("saldo_mes", 0))
                saldo_ac = f.get("saldo_acumulado", 0)
                alerta   = saldo_ac < 0

                fill = C_ROJO_CLR if alerta else (C_VERDE_CLR if i % 2 == 0 else C_BLANCO)
                estado = "⚠ Déficit" if alerta else "✓ Ok"

                vals = [f.get("mes"), acta, amort, neto_ing, costos, flujo_n, saldo_ac, estado]
                fmts = [None, FMT_PESOS, FMT_PESOS, FMT_PESOS, FMT_PESOS, FMT_PESOS, FMT_PESOS, None]
                for col_n, (val, fmt) in enumerate(zip(vals, fmts), 1):
                    _data_cell(ws5, i, col_n, val, fmt=fmt, fill=fill,
                               align="right" if col_n > 1 else "center",
                               bold=(col_n == 7), color="CC0000" if (alerta and col_n == 7) else "000000")

            # Capital de trabajo
            cap_row = len(flujo_data) + 4
            cap = body.flujo.get("capital_trabajo_adicional", 0)
            anticipo_rec = body.flujo.get("anticipo_recibido", 0)
            ws5.merge_cells(f"A{cap_row}:D{cap_row}")
            _data_cell(ws5, cap_row, 1, f"Anticipo recibido al inicio:", bold=True, fill=C_VERDE_CLR)
            _data_cell(ws5, cap_row, 5, anticipo_rec, fmt=FMT_PESOS, bold=True, fill=C_VERDE_CLR, align="right")
            ws5.merge_cells(f"A{cap_row+1}:D{cap_row+1}")
            _data_cell(ws5, cap_row+1, 1, "Capital de trabajo adicional requerido:", bold=True,
                       fill=C_ROJO_CLR if cap > 0 else C_VERDE_CLR)
            _data_cell(ws5, cap_row+1, 5, cap, fmt=FMT_PESOS, bold=True,
                       fill=C_ROJO_CLR if cap > 0 else C_VERDE_CLR, align="right")

            # Gráfico de barras — flujo neto mensual
            if len(flujo_data) >= 2:
                bar = BarChart()
                bar.type   = "col"
                bar.title  = "Flujo neto mensual"
                bar.y_axis.title = "COP"
                bar.x_axis.title = "Mes"
                bar.width  = 18
                bar.height = 10
                data_bar = Reference(ws5, min_col=6, min_row=2, max_row=2 + len(flujo_data))
                cats_bar = Reference(ws5, min_col=1, min_row=3, max_row=2 + len(flujo_data))
                bar.add_data(data_bar, titles_from_data=True)
                bar.set_categories(cats_bar)
                ws5.add_chart(bar, "A" + str(cap_row + 4))
        else:
            ws5.cell(row=1, column=1, value="No hay datos de flujo de caja. Ejecute el flujo desde la calculadora primero.")

        for col, w in zip("ABCDEFGH", [8, 16, 16, 16, 16, 16, 16, 10]):
            ws5.column_dimensions[col].width = w
        _freeze(ws5)

        # ── Guardar y retornar ────────────────────────────────────────────
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{n_fname}"'},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generando Excel: {e}")
